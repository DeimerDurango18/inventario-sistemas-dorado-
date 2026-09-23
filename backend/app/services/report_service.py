"""Generación de reportes Excel con openpyxl."""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.asset import Activo
from app.models.catalog import Proveedor
from app.models.documental import Acta, Baja, Garantia
from app.models.loan import Prestamo
from app.models.maintenance import Mantenimiento
from app.models.movement import Movimiento
from app.models.stock import MovimientoStock, StockItem
from app.models.system import Ticket

_HEADER_FILL = PatternFill("solid", fgColor="696CFF")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _estilo(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"


def _fecha(v):
    if not v:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v)


def _moneda(v):
    if v is None:
        return ""
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return v


def _estado_garantia(g: Garantia) -> str:
    hoy = date.today()
    fin = g.fin.date() if isinstance(g.fin, datetime) else (g.fin if isinstance(g.fin, date) else None)
    if not fin:
        return "Sin fecha fin"
    dias = (fin - hoy).days
    if dias < 0:
        return "VENCIDA"
    if dias <= 60:
        return "POR VENCER"
    return "VIGENTE"


def reporte_activos(db: Session) -> bytes:
    rows = db.scalars(select(Activo).order_by(Activo.codigo)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    _estilo(
        ws,
        ["Código", "Tipo", "Serial", "Cód. inventario", "Placa", "Categoría", "Marca", "Modelo", "Sede", "Ubicación", "Responsable", "Estado", "Valor", "Factura", "F. compra", "F. garantía", "Proveedor", "Observaciones"],
        [
            [
                a.codigo, a.tipo, a.serial, a.codigo_inventario, a.placa,
                a.categoria.nombre if a.categoria else "", a.marca.nombre if a.marca else "",
                a.modelo.nombre if a.modelo else "",
                a.ubicacion.sede.nombre if a.ubicacion and a.ubicacion.sede else "",
                a.ubicacion.nombre if a.ubicacion else "",
                a.responsable.nombre if a.responsable else "", a.estado.nombre if a.estado else "",
                _moneda(a.valor_adquisicion), a.factura_numero,
                _fecha(a.fecha_adquisicion), _fecha(a.fecha_fin_garantia),
                a.proveedor.nombre if a.proveedor else "", a.observaciones,
            ]
            for a in rows
        ],
        [12, 14, 16, 16, 12, 18, 14, 14, 14, 18, 22, 14, 14, 14, 12, 12, 16, 40],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_movimientos(db: Session) -> bytes:
    rows = db.scalars(select(Movimiento).order_by(Movimiento.fecha.desc())).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    _estilo(
        ws,
        ["Número", "Fecha", "Tipo", "Activo", "Resp. anterior", "Resp. nuevo", "Origen", "Destino", "Motivo", "Observaciones", "Acta", "Estado"],
        [
            [
                m.numero, _fecha(m.fecha), m.tipo, m.activo.codigo if m.activo else m.activo_id,
                m.responsable_anterior.nombre if m.responsable_anterior else "",
                m.responsable_nuevo.nombre if m.responsable_nuevo else "",
                m.origen_ubicacion.nombre if m.origen_ubicacion else "",
                m.destino_ubicacion.nombre if m.destino_ubicacion else "",
                m.motivo, m.observaciones,
                m.acta_id or "", m.estado,
            ]
            for m in rows
        ],
        [12, 16, 14, 12, 20, 20, 18, 18, 20, 40, 12, 12],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_mantenimientos(db: Session) -> bytes:
    rows = db.scalars(select(Mantenimiento).order_by(Mantenimiento.id.desc())).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Mantenimientos"
    _estilo(
        ws,
        ["Número", "Activo", "Tipo", "Programado", "Ejecutado", "Técnico", "Diagnóstico", "Actividades", "Resultado", "Costo", "Proveedor", "Estado"],
        [
            [
                m.numero, m.activo.codigo if m.activo else m.activo_id, m.tipo,
                _fecha(m.fecha_programada), _fecha(m.fecha_ejecucion),
                m.tecnico.nombre if m.tecnico else "", m.diagnostico, m.actividades,
                m.resultado, _moneda(m.costo),
                m.proveedor.nombre if m.proveedor else "", m.estado,
            ]
            for m in rows
        ],
        [12, 12, 14, 16, 16, 22, 30, 30, 30, 12, 16, 12],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_garantias(db: Session) -> bytes:
    rows = db.scalars(select(Garantia).order_by(Garantia.fin)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Garantías"
    _estilo(
        ws,
        ["Número activo", "Fabricante", "Proveedor", "Fecha compra", "Inicio", "Fin", "Estado", "Condiciones"],
        [
            [
                g.activo.codigo if g.activo else g.activo_id, g.fabricante,
                g.proveedor.nombre if g.proveedor else "", _fecha(g.fecha_compra),
                _fecha(g.inicio), _fecha(g.fin), _estado_garantia(g), g.condiciones,
            ]
            for g in rows
        ],
        [12, 16, 18, 12, 12, 12, 14, 40],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_tickets(db: Session) -> bytes:
    rows = db.scalars(select(Ticket).order_by(Ticket.id.desc())).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    _estilo(
        ws,
        ["Número", "Fecha", "Cliente", "Categoría", "Prioridad", "Descripción", "Solución", "F. solución", "Estado"],
        [
            [t.numero, _fecha(t.fecha), t.solicitante, t.categoria, t.prioridad, t.descripcion, t.solucion, _fecha(t.fecha_solucion), t.estado]
            for t in rows
        ],
        [12, 16, 22, 16, 10, 40, 40, 16, 12],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_bajas(db: Session) -> bytes:
    rows = db.scalars(select(Baja).order_by(Baja.id.desc())).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Bajas"
    _estilo(
        ws,
        ["Número", "Activo", "Motivo tipo", "Descripción motivo", "Estado físico", "F. solicitud", "F. aprobación", "Estado"],
        [
            [
                b.numero, b.activo.codigo if b.activo else b.activo_id, b.motivo_tipo,
                b.motivo_descripcion, b.estado_fisico, _fecha(b.fecha_solicitud),
                _fecha(b.fecha_aprobacion), b.estado,
            ]
            for b in rows
        ],
        [12, 12, 14, 40, 24, 16, 16, 12],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_stock(db: Session) -> bytes:
    items = db.scalars(select(StockItem).order_by(StockItem.nombre)).all()
    movs = db.scalars(select(MovimientoStock).order_by(MovimientoStock.fecha.desc())).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Existencias"
    _estilo(
        ws,
        ["Código", "Ítem", "Tipo", "Marca", "Modelo", "Categoría", "Stock actual", "Stock mínimo", "Por ubicación", "Valor unitario", "Valor total", "Estado"],
        [
            [
                i.codigo or "", i.nombre, i.tipo,
                i.marca.nombre if i.marca else "", i.modelo.nombre if i.modelo else "",
                i.categoria.nombre if i.categoria else "",
                i.cantidad_stock, i.stock_minimo,
                ", ".join(f"{u.ubicacion.nombre}: {u.cantidad}" for u in i.ubicaciones if u.cantidad) or "—",
                _moneda(i.valor_unitario),
                _moneda((i.valor_unitario or 0) * i.cantidad_stock),
                "Activo" if i.activo else "Inactivo",
            ]
            for i in items
        ],
        [12, 30, 12, 16, 16, 18, 12, 12, 34, 14, 14, 10],
    )
    ws2 = wb.create_sheet("Movimientos")
    _estilo(
        ws2,
        ["Número", "Fecha", "Tipo", "Referencia", "Cantidad", "Ubicación", "Documento", "Destino", "Valor", "Ajuste (antes→nuevo)", "Usuario", "Estado", "Acta"],
        [
            [
                m.numero, _fecha(m.fecha), m.tipo,
                (m.item.nombre if m.item else m.activo.codigo if m.activo else ""),
                m.cantidad, m.ubicacion.nombre if m.ubicacion else "",
                m.documento or "", m.destino or "", _moneda(m.valor),
                f"{m.stock_anterior} → {m.nuevo_stock}" if m.tipo == "AJUSTE" else "",
                m.usuario.username if m.usuario else "", m.estado, m.acta_id or "",
            ]
            for m in movs
        ],
        [12, 16, 12, 30, 10, 18, 14, 20, 14, 18, 16, 12, 10],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_GENERATORS = {
    "activos": reporte_activos,
    "movimientos": reporte_movimientos,
    "mantenimientos": reporte_mantenimientos,
    "garantias": reporte_garantias,
    "tickets": reporte_tickets,
    "bajas": reporte_bajas,
    "stock": reporte_stock,
}


def generar(nombre: str, db: Session) -> bytes | None:
    fn = _GENERATORS.get(nombre)
    if not fn:
        return None
    return fn(db)